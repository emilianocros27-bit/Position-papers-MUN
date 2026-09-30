#!/usr/bin/env python3
"""Build a polished MUN position paper from validated JSON input."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlparse

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - exercised only without dependency
    raise SystemExit(
        "Missing dependency: python-docx. Use the environment's document runtime "
        "or install plugins/mun-position-paper/requirements.txt."
    ) from exc


REQUIRED_TEXT = ("committee", "delegation", "topic")
REQUIRED_LISTS = ("background", "national_position", "solutions")
PLACEHOLDER_RE = re.compile(
    r"\[(?:pa[ií]s|delegaci[oó]n|comit[eé]|t[oó]pico|nombre|escuela|fuente|url)\]",
    re.IGNORECASE,
)


class InputError(ValueError):
    """Raised when a position-paper input file is incomplete or malformed."""


def _nonempty_text(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip())


def validate_payload(data: object) -> dict:
    if not isinstance(data, dict):
        raise InputError("The JSON root must be an object.")

    errors: list[str] = []
    for key in REQUIRED_TEXT:
        if not _nonempty_text(data.get(key)):
            errors.append(f"'{key}' must be a non-empty string")

    for key in REQUIRED_LISTS:
        value = data.get(key)
        if not isinstance(value, list) or not value:
            errors.append(f"'{key}' must be a non-empty list")
        elif any(not _nonempty_text(item) for item in value):
            errors.append(f"every item in '{key}' must be a non-empty string")

    bibliography = data.get("bibliography", [])
    if not isinstance(bibliography, list):
        errors.append("'bibliography' must be a list")
    else:
        for index, entry in enumerate(bibliography, start=1):
            if not isinstance(entry, dict):
                errors.append(f"bibliography entry {index} must be an object")
                continue
            for key in ("author", "title", "url"):
                if not _nonempty_text(entry.get(key)):
                    errors.append(f"bibliography entry {index} requires '{key}'")
            url = entry.get("url")
            if _nonempty_text(url):
                parsed = urlparse(url)
                if parsed.scheme not in {"http", "https"} or not parsed.netloc:
                    errors.append(f"bibliography entry {index} has an invalid URL")

    if not data.get("draft", False) and not bibliography:
        errors.append("a final paper requires at least one bibliography entry")

    flag_path = data.get("flag_path")
    if flag_path is not None:
        if not _nonempty_text(flag_path):
            errors.append("'flag_path' must be a non-empty path when present")
        elif not Path(flag_path).expanduser().is_file():
            errors.append(f"flag image not found: {flag_path}")

    if not data.get("draft", False):
        serialized = json.dumps(data, ensure_ascii=False)
        if PLACEHOLDER_RE.search(serialized) or "TODO" in serialized.upper():
            errors.append("final input contains an unfinished placeholder")

    if errors:
        raise InputError("Invalid input:\n- " + "\n- ".join(errors))
    return data


def _set_font(run, size: float = 11.0, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def _set_cell_margins(cell, top: int = 40, start: int = 40, bottom: int = 40, end: int = 40) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def _set_repeat_table_layout_fixed(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_font(run, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def _add_metadata_line(paragraph, label: str, value: str) -> None:
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    label_run = paragraph.add_run(f"{label}: ")
    _set_font(label_run, 11, bold=True)
    value_run = paragraph.add_run(value)
    _set_font(value_run, 11)


def _add_body_paragraph(document, text: str, *, keep_with_next: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text.strip())
    _set_font(run, 11)


def _add_external_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    for item in (fonts, color, underline, size):
        properties.append(item)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_document(data: dict, output_path: Path) -> None:
    document = Document()
    author = data.get("delegate", "").strip() if isinstance(data.get("delegate"), str) else ""
    document.core_properties.author = author
    document.core_properties.last_modified_by = author
    document.core_properties.title = f"{data['delegation']} - {data['topic']}"
    document.core_properties.subject = "MUN position paper"
    document.core_properties.comments = ""
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _remove_table_borders(table)
    _set_repeat_table_layout_fixed(table)
    left, right = table.rows[0].cells
    left.width = Inches(5.55)
    right.width = Inches(1.25)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(left, 0, 0, 0, 80)
    _set_cell_margins(right, 0, 80, 0, 0)

    metadata = [
        ("Comité", data["committee"]),
        ("Delegación", data["delegation"]),
    ]
    if _nonempty_text(data.get("delegate")):
        metadata.append(("Delegado", data["delegate"].strip()))
    if _nonempty_text(data.get("school")):
        metadata.append(("Escuela", data["school"].strip()))
    if _nonempty_text(data.get("conference")):
        metadata.append(("Conferencia", data["conference"].strip()))
    metadata.append(("Tópico", data["topic"]))

    left.text = ""
    for index, (label, value) in enumerate(metadata):
        paragraph = left.paragraphs[0] if index == 0 else left.add_paragraph()
        _add_metadata_line(paragraph, label, value)

    right.text = ""
    if data.get("flag_path"):
        flag_paragraph = right.paragraphs[0]
        flag_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        flag_paragraph.add_run().add_picture(
            str(Path(data["flag_path"]).expanduser().resolve()), width=Inches(1.25)
        )

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)

    if data.get("draft", False):
        draft = document.add_paragraph()
        draft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        draft.paragraph_format.space_after = Pt(8)
        run = draft.add_run("BORRADOR DE DEMOSTRACIÓN - NO ENTREGAR")
        _set_font(run, 10, bold=True)

    if _nonempty_text(data.get("opening")):
        _add_body_paragraph(document, data["opening"])

    show_headings = bool(data.get("show_section_headings", False))
    if show_headings:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(5)
        heading.paragraph_format.space_after = Pt(3)
        _set_keep_with_next(heading)
        _set_font(heading.add_run("Contexto y acción previa"), 11, bold=True)
    for paragraph in data["background"]:
        _add_body_paragraph(document, paragraph)

    if show_headings:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(5)
        heading.paragraph_format.space_after = Pt(3)
        _set_keep_with_next(heading)
        _set_font(heading.add_run("Postura de la delegación"), 11, bold=True)
    for paragraph in data["national_position"]:
        _add_body_paragraph(document, paragraph)

    solutions_heading = document.add_paragraph()
    solutions_heading.paragraph_format.space_before = Pt(5)
    solutions_heading.paragraph_format.space_after = Pt(3)
    _set_keep_with_next(solutions_heading)
    _set_font(solutions_heading.add_run("Soluciones"), 11, bold=True)

    for solution in data["solutions"]:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.17)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.05
        _set_font(paragraph.add_run(solution.strip()), 11)

    bibliography = data.get("bibliography", [])
    if bibliography:
        bibliography_heading = document.add_paragraph()
        bibliography_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bibliography_heading.paragraph_format.space_before = Pt(12)
        bibliography_heading.paragraph_format.space_after = Pt(8)
        _set_keep_with_next(bibliography_heading)
        title = data.get("bibliography_title") or "Bibliografía"
        _set_font(bibliography_heading.add_run(title), 13, bold=True)

        for entry in bibliography:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            date = entry.get("date", "s.f.").strip() if isinstance(entry.get("date"), str) else "s.f."
            prefix = f"{entry['author'].strip()}. ({date}). {entry['title'].strip()}. "
            _set_font(paragraph.add_run(prefix), 9.5)
            _add_external_hyperlink(paragraph, entry["url"].strip(), entry["url"].strip())

    if data.get("include_page_numbers", True):
        _add_page_number(section.footer.paragraphs[0])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def convert_to_pdf(docx_path: Path, pdf_output: Path) -> None:
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if office is None:
        raise RuntimeError("LibreOffice/soffice was not found; the DOCX was created successfully.")
    pdf_output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mun-position-paper-") as temp_dir:
        command = [
            office,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            temp_dir,
            str(docx_path.resolve()),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, check=False)
        converted = Path(temp_dir) / f"{docx_path.stem}.pdf"
        if completed.returncode != 0 or not converted.is_file():
            message = completed.stderr.strip() or completed.stdout.strip() or "unknown conversion error"
            raise RuntimeError(f"PDF conversion failed: {message}")
        shutil.copy2(converted, pdf_output)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="UTF-8 JSON input file")
    parser.add_argument("--output", type=Path, default=Path("position-paper.docx"))
    parser.add_argument("--check-only", action="store_true", help="validate without writing files")
    parser.add_argument("--pdf", action="store_true", help="also convert the DOCX to PDF")
    parser.add_argument("--pdf-output", type=Path, help="optional PDF destination")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        with args.input.open("r", encoding="utf-8") as handle:
            data = validate_payload(json.load(handle))
        if args.check_only:
            print(json.dumps({"valid": True, "input": str(args.input)}, ensure_ascii=False))
            return 0
        build_document(data, args.output)
        result: dict[str, object] = {"valid": True, "docx": str(args.output)}
        if args.pdf:
            pdf_output = args.pdf_output or args.output.with_suffix(".pdf")
            convert_to_pdf(args.output, pdf_output)
            result["pdf"] = str(pdf_output)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except (OSError, json.JSONDecodeError, InputError, RuntimeError) as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
